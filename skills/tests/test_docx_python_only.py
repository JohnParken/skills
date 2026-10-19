"""Exercise documented DOCX workflows with no executable search path."""

import importlib.util
import os
import re
import stat
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "docx" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from office.package import pack_document, unpack_document

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class PythonOnlyDocxTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="docx-python-only-test-")
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name)
        self.env = dict(os.environ, PATH="", PYTHONDONTWRITEBYTECODE="1")
        instructions = (ROOT / "docx" / "SKILL.md").read_text(encoding="utf-8")
        example = re.search(r"```python\n(.*?)```", instructions, re.S).group(1)
        result = subprocess.run(
            [sys.executable, "-B", "-c", example], cwd=self.work,
            env=self.env, capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.original = self.work / "output.docx"

    def cli(self, script, *args, expected=0):
        result = subprocess.run(
            [sys.executable, "-B", str(SCRIPTS / script), *map(str, args)],
            cwd=self.work, env=self.env, capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        return result

    def test_documented_generation_editing_reading_and_validation(self):
        unpacked = self.work / "unpacked"
        self.cli("office/package.py", "unpack", self.original, unpacked)
        self.cli("merge_runs.py", unpacked)
        xml = unpacked / "word/document.xml"
        tree = etree.parse(str(xml))
        for node in tree.iter(f"{{{W}}}t"):
            if node.text == "The team completed the first milestone.":
                node.text = "团队已完成第二个里程碑。"
        tree.write(str(xml), xml_declaration=True, encoding="UTF-8")
        edited = self.work / "edited.docx"
        self.cli("office/package.py", "pack", unpacked, edited)
        validation = self.cli("office/validate.py", edited)
        self.assertIn("All validations PASSED!", validation.stdout)
        text_file = self.work / "content.txt"
        self.cli("read_docx.py", edited, "--include-related", "-o", text_file)
        text = text_file.read_text(encoding="utf-8")
        self.assertIn("团队已完成第二个里程碑。", text)
        self.assertIn("Milestone\tStatus", text)
        self.assertNotIn("The team completed the first milestone.", text)
        reopened = Document(edited)
        self.assertEqual(len(reopened.tables), 1)
        with zipfile.ZipFile(self.original) as before, zipfile.ZipFile(edited) as after:
            self.assertIsNone(after.testzip())
            self.assertEqual(set(before.namelist()), set(after.namelist()))
            for name in before.namelist():
                if name != "word/document.xml":
                    self.assertEqual(before.read(name), after.read(name), name)

    def test_comments_remain_structurally_valid_when_anchored(self):
        annotated = self.work / "annotated.docx"
        self.cli("comment.py", self.original, "核实费用 & 范围", "-o", annotated)
        unpacked = unpack_document(annotated, self.work / "annotated")
        comment = etree.parse(str(unpacked / "word/comments.xml")).getroot()[0]
        comment_id = comment.get(f"{{{W}}}id")
        xml = unpacked / "word/document.xml"
        tree = etree.parse(str(xml))
        paragraph = next(tree.iter(f"{{{W}}}p"))
        start = etree.Element(f"{{{W}}}commentRangeStart", {f"{{{W}}}id": comment_id})
        paragraph.insert(1 if paragraph.find(f"{{{W}}}pPr") is not None else 0, start)
        paragraph.append(etree.Element(f"{{{W}}}commentRangeEnd", {f"{{{W}}}id": comment_id}))
        run = etree.SubElement(paragraph, f"{{{W}}}r")
        etree.SubElement(run, f"{{{W}}}commentReference", {f"{{{W}}}id": comment_id})
        tree.write(str(xml), xml_declaration=True, encoding="UTF-8")
        final = pack_document(unpacked, self.work / "anchored.docx")
        self.cli("office/validate.py", final)
        extracted = self.cli("read_docx.py", final, "--include-related").stdout
        self.assertIn("核实费用 & 范围", extracted)

    @unittest.skipUnless(importlib.util.find_spec("docxtpl"), "optional docxtpl not installed")
    def test_optional_template_filling_and_missing_data(self):
        from docxtpl import DocxTemplate
        from jinja2 import Environment, StrictUndefined, UndefinedError

        template = self.work / "template.docx"
        document = Document(self.original)
        document.paragraphs[0].text = "{{ company }}"
        document.save(template)
        env = Environment(undefined=StrictUndefined)
        filled = DocxTemplate(template)
        filled.render({"company": "ACME & 中文"}, jinja_env=env, autoescape=True)
        output = self.work / "filled.docx"
        filled.save(output)
        self.cli("office/validate.py", output)
        self.assertIn("ACME & 中文", self.cli("read_docx.py", output).stdout)
        with self.assertRaises(UndefinedError):
            DocxTemplate(template).render({}, jinja_env=env, autoescape=True)

    def test_validation_detects_missing_relationship_target(self):
        unpacked = unpack_document(self.original, self.work / "broken")
        rels = unpacked / "word/_rels/document.xml.rels"
        rels.write_text(rels.read_text().replace('Target="styles.xml"', 'Target="missing-styles.xml"'))
        broken = pack_document(unpacked, self.work / "broken.docx")
        result = self.cli("office/validate.py", broken, expected=1)
        self.assertIn("missing-styles.xml", result.stdout + result.stderr)

    def test_redlining_cli_without_git(self):
        unpacked = unpack_document(self.original, self.work / "redlined")
        xml = unpacked / "word/document.xml"
        tree = etree.parse(str(xml))
        first_text = next(tree.iter(f"{{{W}}}t"))
        first_text.text = "未经追踪的修改"
        tree.write(str(xml), xml_declaration=True, encoding="UTF-8")
        untracked = pack_document(unpacked, self.work / "untracked.docx")
        failed = self.cli(
            "office/validate.py", untracked, "--original", self.original,
            "--author", "Editor", expected=1,
        )
        self.assertIn("--- original", failed.stdout)
        self.assertIn("+未经追踪的修改", failed.stdout)

        unpacked = unpack_document(self.original, self.work / "tracked-parts")
        xml = unpacked / "word/document.xml"
        tree = etree.parse(str(xml))
        paragraph = next(tree.iter(f"{{{W}}}p"))
        insertion = etree.SubElement(paragraph, f"{{{W}}}ins", {
            f"{{{W}}}id": "0", f"{{{W}}}author": "Editor",
            f"{{{W}}}date": "2026-01-01T00:00:00Z",
        })
        run = etree.SubElement(insertion, f"{{{W}}}r")
        etree.SubElement(run, f"{{{W}}}t").text = "已复核"
        tree.write(str(xml), xml_declaration=True, encoding="UTF-8")
        tracked = pack_document(unpacked, self.work / "tracked.docx")
        self.cli("office/validate.py", tracked, "--original", self.original, "--author", "Editor")

    def test_unsafe_archives_do_not_replace_destination(self):
        for name, symlink in (("../escape.txt", False), ("link", True)):
            with self.subTest(name=name):
                bad = self.work / "unsafe.docx"
                with zipfile.ZipFile(self.original) as original, zipfile.ZipFile(bad, "w") as archive:
                    for item in original.infolist():
                        archive.writestr(item, original.read(item.filename))
                    item = zipfile.ZipInfo(name)
                    if symlink:
                        item.create_system = 3
                        item.external_attr = (stat.S_IFLNK | 0o777) << 16
                    archive.writestr(item, "target")
                dest = self.work / "unsafe-output"
                with self.assertRaises(ValueError):
                    unpack_document(bad, dest)
                self.assertFalse(dest.exists())
                self.assertFalse((self.work / "escape.txt").exists())

    def test_pack_boundaries_and_existing_outputs(self):
        unpacked = unpack_document(self.original, self.work / "parts")
        with self.assertRaisesRegex(ValueError, "outside"):
            pack_document(unpacked, unpacked / "recursive.docx")
        (unpacked / "linked.xml").symlink_to(unpacked / "word/document.xml")
        before = self.original.read_bytes()
        with self.assertRaisesRegex(ValueError, "links"):
            pack_document(unpacked, self.original)
        self.assertEqual(before, self.original.read_bytes())
        with self.assertRaisesRegex(ValueError, "empty"):
            unpack_document(self.original, unpacked)

    def test_invalid_zip_and_missing_parts_fail(self):
        invalid = self.work / "invalid.docx"
        invalid.write_bytes(b"not a ZIP")
        self.cli("office/package.py", "unpack", invalid, self.work / "invalid", expected=1)
        with zipfile.ZipFile(invalid, "w") as archive:
            archive.writestr("word/document.xml", "<document/>")
        with self.assertRaisesRegex(ValueError, "Missing required"):
            unpack_document(invalid, self.work / "missing")

    def test_dotx_package_roundtrip(self):
        unpacked = unpack_document(self.original, self.work / "template-parts")
        types = unpacked / "[Content_Types].xml"
        types.write_text(types.read_text().replace(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
        ))
        template = pack_document(unpacked, self.work / "reference.dotx")
        roundtrip = unpack_document(template, self.work / "template-roundtrip")
        output = pack_document(roundtrip, self.work / "copy.dotx")
        self.cli("office/validate.py", output)
        self.assertIn("Project update", self.cli("read_docx.py", output).stdout)


if __name__ == "__main__":
    unittest.main()
